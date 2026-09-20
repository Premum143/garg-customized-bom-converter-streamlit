from __future__ import annotations

from collections import OrderedDict
from io import BytesIO
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


FG_HEADERS = [
    "Sl_No",
    "FG Item ID",
    "FG Item Name",
    "FG UOM",
    "BOM Number",
    "BOM Name",
    "FG Store",
    "RM Store",
    "Scrap/By-product Store",
    "WIP Store",
    "BOM Description",
    "FG Cost Allocation",
    "FG Comment",
    "Comment",
]

RM_HEADERS = [
    "Sl_No",
    "FG Item ID",
    "BOM Number",
    "#",
    "Item Id",
    "Item Description",
    "Quantity",
    "Unit",
    "Comment",
]

ROUTING_HEADERS = [
    "Sl_No",
    "FG Item ID",
    "BOM Number",
    "Routing Item Type",
    "RM/FG Sl_No",
    "Routing Item ID",
    "#",
    "Routing Number",
    "Routing Name",
    "Input Variant",
    "Output Variant",
    "Est. Time / Unit (sec)",
    "Workstation Group",
    "Workstation",
    "Operator",
    "Require Prev Routing Done",
    "Comment",
]

MASTER_HEADERS = [
    "Item ID",
    "Item Name",
    "Product/Service",
    "Item Type (Buy/Sell/Both)",
    "Unit of Measurement",
    "HSN Code",
    "Item Category",
    "Default Price",
    "Regular Buying Price",
    "Wholesale Buying Price",
    "Regular Selling Price",
    "MRP",
    "Dealer Price",
    "Distributor Price",
    "Current Stock",
    "Min Stock Level",
    "Max Stock Level",
    "Tax",
]

ROUTING_MAP = [
    ("CUTTING", "CU-01", "Cutting"),
    ("CHAMFER", "CH-01", "Chamfering"),
    ("BENDING", "BE-01", "Bending"),
    ("ROLLING", "RO-01", "Rolling"),
    ("PREMACHINING", "PR-01", "Pre-machining"),
    ("DRILLING", "DR-01", "Drilling"),
    ("FABRICATION", "FA-01", "Fabrication"),
]

DEFAULT_BOM_NUMBER = "GARG (Customized Assembly BOM)"
DEFAULT_STORES = {
    "fg_store": "Garg Unit - 1",
    "rm_store": "Garg Unit - 2",
    "scrap_store": "Default Reject Store",
    "wip_store": "Default WIP Store",
}


class ConversionError(Exception):
    pass


def clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def normalize_label(text):
    return re.sub(r"[^a-z0-9]+", "", clean(text).lower())


def parse_metadata(ws):
    metadata = {}
    for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
        values = [clean(value) for value in row]
        upper_values = [value.upper() for value in values]
        if "DRAWING NO" in upper_values and any(value in {"SL.NO", "SL NO"} for value in upper_values):
            break
        for index, value in enumerate(values[:-1]):
            key = normalize_label(value)
            next_value = values[index + 1]
            if not next_value:
                continue
            if key in {"description", "discription"}:
                metadata["description"] = next_value
            elif key == "customer":
                metadata["customer"] = next_value
            elif key in {"assemblydrawingno", "assydrgno"}:
                metadata["fg_code"] = next_value
            elif key in {"refdrawingno", "refdrgno"}:
                metadata["ref_drawing"] = next_value
            elif key == "pono":
                metadata["po_no"] = next_value
            elif key == "qty":
                metadata["fg_qty_label"] = next_value
    return metadata


def find_bom_header(ws):
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        values = [clean(value) for value in row]
        upper_values = [value.upper() for value in values]
        if "DRAWING NO" in upper_values and any(value in {"SL.NO", "SL NO"} for value in upper_values):
            return row_no, values
    raise ConversionError(f"No BOM header found in sheet {ws.title}")


def component_row_dict(headers, values):
    row = {}
    for index, header in enumerate(headers):
        if not header:
            continue
        row[header.upper()] = clean(values[index]) if index < len(values) else ""
    return row


def is_invalid_marker(value):
    return clean(value).upper() in {"", "N/A", "NA", "NIA", "-", "--", "---"}


def parse_quantity(value, default=None):
    text = clean(value)
    if not text:
        return default
    try:
        number = float(text)
    except ValueError:
        return default
    if number.is_integer():
        return int(number)
    return number


def is_component_row(row):
    drawing_no = clean(row.get("DRAWING NO"))
    material_code = clean(row.get("MATERIAL CODE"))
    qty = parse_quantity(row.get("QTY") or row.get("QTY/SET"))
    return bool(drawing_no and material_code and qty not in {None, 0})


def format_size(row):
    parts = []
    for key in ("THK", "THICKNESS", "WIDTH", "LENGTH"):
        value = clean(row.get(key))
        if value and value not in {"0", "0.0"}:
            parts.append(value)
    return " x ".join(parts)


def build_item_name(row):
    description = clean(row.get("DESCRIPTION") or row.get("DISCRIPTION"))
    material_code = clean(row.get("MATERIAL CODE"))
    size = format_size(row)
    parts = [part for part in (description, material_code, size) if part]
    return " - ".join(parts)


def component_category(row):
    material_code = clean(row.get("MATERIAL CODE")).upper()
    description = clean(row.get("DESCRIPTION") or row.get("DISCRIPTION")).upper()
    if material_code.startswith(("PL", "SH")) or "PLATE" in description or "SHEET" in description:
        return "Plate Component"
    if material_code.startswith(("CH", "IS", "UC", "UB", "RHS", "SHS", "PIPE", "RD", "ISA")):
        return "Structural Component"
    if material_code.startswith(("BOP", "BRG")) or any(
        word in description for word in ("BOLT", "NUT", "BEARING", "WASHER", "SCREW")
    ):
        return "Bought Out Part"
    return "Customized Component"


def add_master(master_items, item_id, item_name, item_type, uom, category):
    if not item_id or item_id in master_items:
        return
    master_items[item_id] = [
        item_id,
        item_name,
        "Product",
        item_type,
        uom,
        "",
        category,
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
    ]


def make_bom_description(metadata):
    details = []
    if metadata.get("customer"):
        details.append(f"Customer: {metadata['customer']}")
    if metadata.get("po_no"):
        details.append(f"PO: {metadata['po_no']}")
    return " | ".join(details)


def parse_section_header(raw_value, fg_code):
    text = clean(raw_value)
    if not text:
        return None
    code_match = re.match(r"^([A-Za-z0-9_-]+?A\d+)[ _-]*(.*)$", text)
    if not code_match:
        return None
    code = code_match.group(1)
    name = code_match.group(2).strip(" _-")
    fg_match = re.match(r"^(.*?[_-]A)\d+$", fg_code)
    if fg_match and not code.startswith(fg_match.group(1)):
        return None
    if code == fg_code:
        return None
    return {
        "code": code,
        "name": name.replace("_", " ") if name else code,
        "raw": text,
    }


def find_section_header(row, fg_code):
    for key in ("SL.NO", "SL NO", "ITEM NO.", "ITEM NO", "DRAWING NO"):
        section = parse_section_header(row.get(key), fg_code)
        if section:
            return section
    return None


def should_record_skipped_row(row):
    preview = (
        clean(row.get("DRAWING NO")),
        clean(row.get("DESCRIPTION") or row.get("DISCRIPTION")),
        clean(row.get("MATERIAL CODE")),
        clean(row.get("QTY") or row.get("QTY/SET")),
    )
    return any(preview)


def parse_customized_bom(source, filename=None):
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if "BOM" in workbook.sheetnames:
            ws = workbook["BOM"]
        elif workbook.sheetnames and workbook.sheetnames[0] == "Order Voucher":
            raise ConversionError("Tally Order Voucher export detected; this tool supports BOM-format workbooks only.")
        else:
            ws = workbook[workbook.sheetnames[0]]

        metadata = parse_metadata(ws)
        header_row, headers = find_bom_header(ws)
        fg_code = metadata.get("fg_code")
        if not fg_code:
            raise ConversionError(f"Could not determine FG code for {filename or 'uploaded file'}")
        if not metadata.get("description"):
            metadata["description"] = fg_code

        root_components = []
        sections = []
        current_section = None
        skipped_rows = []

        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            row = component_row_dict(headers, values)
            if not any(row.values()):
                continue

            section = find_section_header(row, fg_code)
            if section:
                current_section = {**section, "rows": []}
                sections.append(current_section)
                continue

            if is_component_row(row):
                if current_section is not None:
                    current_section["rows"].append(row)
                else:
                    root_components.append(row)
                continue

            if should_record_skipped_row(row):
                preview = (
                    clean(row.get("DRAWING NO")),
                    clean(row.get("DESCRIPTION") or row.get("DISCRIPTION")),
                    clean(row.get("MATERIAL CODE")),
                    clean(row.get("QTY") or row.get("QTY/SET")),
                )
                skipped_rows.append(preview)

        return {
            "path": Path(filename) if filename else Path(getattr(source, "name", "uploaded.xlsx")),
            "metadata": metadata,
            "header_row": header_row,
            "mode": "multi_level" if sections else "flat",
            "root_components": root_components,
            "sections": sections,
            "skipped_rows": skipped_rows,
        }
    finally:
        workbook.close()


def component_comment(row):
    bits = []
    boi = clean(row.get("BOI"))
    structure = clean(row.get("STRUCTURE (Y/N)") or row.get("STRUCTURE(Y/N)"))
    if boi and boi.upper() not in {"N/A", "NA"}:
        bits.append(f"BOI: {boi}")
    if structure:
        bits.append(f"Structure: {structure}")
    return " | ".join(bits)


def append_component_rows(fg_sl_no, fg_code, rows, rm_rows, routing_rows, master_items, bom_number):
    for index, row in enumerate(rows, start=1):
        item_id = clean(row.get("DRAWING NO"))
        item_name = build_item_name(row)
        qty = parse_quantity(row.get("QTY") or row.get("QTY/SET"), default=0)
        rm_rows.append([
            fg_sl_no,
            fg_code,
            bom_number,
            index,
            item_id,
            item_name,
            qty,
            "Nos",
            component_comment(row),
        ])
        add_master(master_items, item_id, item_name, "Both", "Nos", component_category(row))

        route_no = 1
        for column, routing_number, routing_name in ROUTING_MAP:
            raw_value = clean(row.get(column))
            if is_invalid_marker(raw_value):
                continue
            routing_rows.append([
                fg_sl_no,
                fg_code,
                bom_number,
                "RM",
                index,
                item_id,
                route_no,
                routing_number,
                routing_name,
                "",
                "",
                "",
                "",
                "",
                "",
                "No",
                raw_value,
            ])
            route_no += 1


def build_tranzact_rows(parsed_products, bom_number=DEFAULT_BOM_NUMBER, stores=None):
    stores = {**DEFAULT_STORES, **(stores or {})}
    fg_rows = []
    rm_rows = []
    routing_rows = []
    master_items = OrderedDict()
    skipped_rows = []
    fg_sl_no = 1
    seen_codes = set()

    for parsed in parsed_products:
        metadata = parsed["metadata"]
        fg_code = metadata["fg_code"]
        if fg_code in seen_codes:
            raise ConversionError(f"Duplicate FG code detected: {fg_code}")
        seen_codes.add(fg_code)

        fg_name = metadata["description"]
        bom_description = make_bom_description(metadata)
        top_fg_sl_no = fg_sl_no

        fg_rows.append([
            top_fg_sl_no,
            fg_code,
            fg_name,
            "Nos",
            bom_number,
            fg_name,
            stores["fg_store"],
            stores["rm_store"],
            stores["scrap_store"],
            stores["wip_store"],
            bom_description,
            100,
            f"Source file: {parsed['path'].name}",
            "",
        ])
        add_master(master_items, fg_code, fg_name, "Sell", "Nos", "Finished Good")
        fg_sl_no += 1
        skipped_rows.extend([(parsed["path"].name, *row) for row in parsed["skipped_rows"]])

        if parsed["mode"] == "multi_level":
            section_no = 1
            for section in parsed["sections"]:
                section_code = section["code"]
                section_name = section["name"]
                sfg_sl_no = fg_sl_no
                fg_rows.append([
                    sfg_sl_no,
                    section_code,
                    section_name,
                    "Nos",
                    bom_number,
                    section_name,
                    stores["fg_store"],
                    stores["rm_store"],
                    stores["scrap_store"],
                    stores["wip_store"],
                    bom_description,
                    100,
                    f"Parent FG: {fg_code}",
                    "",
                ])
                add_master(master_items, section_code, section_name, "Both", "Nos", "Sub Assembly")
                rm_rows.append([
                    top_fg_sl_no,
                    fg_code,
                    bom_number,
                    section_no,
                    section_code,
                    section_name,
                    1,
                    "Nos",
                    "Section-based SFG",
                ])
                append_component_rows(
                    sfg_sl_no,
                    section_code,
                    section["rows"],
                    rm_rows,
                    routing_rows,
                    master_items,
                    bom_number,
                )
                fg_sl_no += 1
                section_no += 1

            if parsed["root_components"]:
                append_component_rows(
                    top_fg_sl_no,
                    fg_code,
                    parsed["root_components"],
                    rm_rows,
                    routing_rows,
                    master_items,
                    bom_number,
                )
        else:
            append_component_rows(
                top_fg_sl_no,
                fg_code,
                parsed["root_components"],
                rm_rows,
                routing_rows,
                master_items,
                bom_number,
            )

    return {
        "fg_rows": fg_rows,
        "rm_rows": rm_rows,
        "routing_rows": routing_rows,
        "master_rows": list(master_items.values()),
        "skipped_rows": skipped_rows,
    }


def write_sheet(workbook, title, headers, rows):
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    autosize(ws)


def build_workbook(result):
    workbook = Workbook()
    del workbook["Sheet"]
    write_sheet(workbook, "FG", FG_HEADERS, result["fg_rows"])
    write_sheet(workbook, "RM", RM_HEADERS, result["rm_rows"])
    write_sheet(workbook, "Routing", ROUTING_HEADERS, result["routing_rows"])
    write_sheet(workbook, "Master File", MASTER_HEADERS, result["master_rows"])
    return workbook


def build_review_workbook(parsed, bom_number=DEFAULT_BOM_NUMBER, stores=None):
    result = build_tranzact_rows([parsed], bom_number=bom_number, stores=stores)
    workbook = build_workbook(result)
    metadata = parsed["metadata"]
    result["fg_code"] = metadata["fg_code"]
    result["fg_name"] = metadata["description"]
    return workbook, result


def autosize(ws):
    for column in ws.columns:
        width = 0
        col_letter = get_column_letter(column[0].column)
        for cell in column[:200]:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value), 80))
        ws.column_dimensions[col_letter].width = max(10, width + 2)


def workbook_to_bytes(workbook):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()
